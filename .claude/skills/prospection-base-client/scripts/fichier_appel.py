"""Étape 9 : fichier d'appel du commercial, une ligne par contact de la liste « entrée », dans l'ordre d'appel.

Colonnes, dans l'ordre de lecture de l'appelant : le contact, l'argument au téléphone et sa source, pourquoi
maintenant, puis le détail (secteur, description, taille, pays, canton ou département, ville, logiciels, signaux sur
la personne et sur l'entreprise, relais à demander, à savoir, standard, LinkedIn, site).

Trois commandes, dans cet ordre :
  briefs     une fiche de faits par entreprise, en lots équilibrés (--lots, 4 par défaut), pour les agents qui
             rédigent secteur, description, taille et l'argument : appels/briefs/L*.json, plus
             appels/briefs/contexte.json (offre, persona, langues) ; réponses attendues dans appels/redaction/L*.json
             (consignes/fichier-appel-redaction.md). Les mêmes briefs servent aux cartes d'appel
             (consignes/cartes-appel.md, puis scripts_appel.py). Si le découpage en lots change, les lots de la
             passe précédente (briefs, rédactions, cartes et sources L*.json) sont rangés dans
             appels/precedent/<date-heure>/ : une ancienne rédaction ne peut plus l'emporter sur la nouvelle.
  sources    des faits numérotés avec leur lien, pour ouvrir chaque argument sur sa source (« J'ai vu ... ») :
             appels/sources/briefs/L*.json ; réponses dans appels/sources/redaction/L*.json
             (consignes/fichier-appel-sources.md). Le lien et la date sont posés par le script, jamais par l'agent.
  assembler  fusionne tout, contrôle ce qui sera dit (tirets, signal négatif ou phrase interdite, lieu d'un autre
             pays, chiffre ou nom absent de la fiche, logiciel nommé sans confirmation) et écrit
             appels/fichier-appel.xlsx (onglets Appels, Relais, Lisez-moi) et appels/fichier-appel.csv, le fichier
             remis au commercial ; avec --maj-ordre, écrit aussi ces colonnes dans le fichier d'ordre d'appel (à
             refaire après chaque ordre_appel.py).

Sources, toutes déjà sur le disque (0 crédit, aucun appel réseau), dans le dossier de campagne :
appels/ordre-appel.csv (ordre_appel.py), prospection-signaux.csv, donnees/signaux/{signaux,recherche,linkedin,
logiciels,entites}.json, apollo/cache/ (réponses d'enrichissement Apollo : parcours des personnes), les relais (export
« - relais » d'exporter_listes.py, sinon apollo/contacts-crees.csv), apollo/contacts-crees.csv (prénoms et noms
inversés par Apollo), apollo/comptes-apollo.json (code postal, standard), les cartes d'appel ; facultatifs : --export
(export CSV de la liste « entrée » depuis l'interface Apollo : effectif, technologies, chiffre d'affaires estimé,
standard) et appels/corrections.json ({id: {pourquoi_maintenant, a_savoir}}, écrit à la main).
Chaque contact est désigné par la colonne contact_id de l'ordre d'appel : id Apollo, ou id local « local-... » d'un
contact absent d'Apollo (ordre_appel.cle_contact). Les rédactions et les sources ne sont reprises que du lot où se
trouve le contact.

Usage (depuis la racine du projet) :
  .venv/bin/python .claude/skills/prospection-base-client/scripts/fichier_appel.py briefs|sources|assembler
      --campagne campagnes/<nom> [--lots 4] [--ordre <csv>] [--relais <csv>] [--export <csv>] [--maj-ordre]
"""
from __future__ import annotations

import argparse
import json
import re
import sys
import unicodedata
from collections import defaultdict
from datetime import date, datetime
from pathlib import Path
from urllib.parse import unquote, urlparse

sys.dont_write_bytecode = True  # aucun __pycache__ dans le dossier du skill, qui se partage tel quel

from commun import ORDRE_RELAIS, Campagne, argument_campagne, charger_localites, code_pays, domaine, ecrire_csv, lire_csv, norm
from ordre_appel import (NOMS_LANGUES, charger_contacts, charger_fiches, cle_contact, comptes_apollo, id_apollo,
                         lire_json, pays_entreprise)
from scripts_appel import (LIEUX, NOMS_PAYS, charger_cartes, interdits_trouves, lieux_a_eviter, lieux_hors_pays,
                           motifs_interdits, pays_surveilles, textes_offre)

REF = date.today()                 # date de référence de la campagne (remplacée au chargement)
NOUVEAU_POSTE_MOIS = 18            # « Nouveau à ce poste » jusqu'à 18 mois
INDICATIF = {"CH": "+41", "FR": "+33", "DE": "+49", "IT": "+39", "AT": "+43", "LI": "+423", "BE": "+32",
             "LU": "+352", "NL": "+31", "ES": "+34", "PT": "+351", "GB": "+44", "IE": "+353", "US": "+1"}
CANTONS = {"AG": "Argovie", "AI": "Appenzell Rhodes-Intérieures", "AR": "Appenzell Rhodes-Extérieures", "BE": "Berne",
           "BL": "Bâle-Campagne", "BS": "Bâle-Ville", "FR": "Fribourg", "GE": "Genève", "GL": "Glaris", "GR": "Grisons",
           "JU": "Jura", "LU": "Lucerne", "NE": "Neuchâtel", "NW": "Nidwald", "OW": "Obwald", "SG": "Saint-Gall",
           "SH": "Schaffhouse", "SO": "Soleure", "SZ": "Schwytz", "TG": "Thurgovie", "TI": "Tessin", "UR": "Uri",
           "VD": "Vaud", "VS": "Valais", "ZG": "Zoug", "ZH": "Zurich"}
# libellé Apollo ou LinkedIn (normalisé) -> code du canton ; l'ordre compte (« basel landschaft » avant « basel »)
CANTON_ALIAS = [("basel landschaft", "BL"), ("basel land", "BL"), ("basel", "BS"), ("bale campagne", "BL"),
                ("bale", "BS"), ("geneva", "GE"), ("geneve", "GE"), ("vaud", "VD"), ("neuchatel", "NE"),
                ("jura", "JU"), ("fribourg", "FR"), ("freiburg", "FR"), ("valais", "VS"), ("wallis", "VS"),
                ("bern", "BE"), ("zurich", "ZH"), ("aargau", "AG"), ("argovie", "AG"), ("lucerne", "LU"),
                ("luzern", "LU"), ("zug", "ZG"), ("zoug", "ZG"), ("solothurn", "SO"), ("soleure", "SO"),
                ("thurgau", "TG"), ("thurgovie", "TG"), ("schwyz", "SZ"), ("st gallen", "SG"), ("saint gall", "SG"),
                ("ticino", "TI"), ("tessin", "TI"), ("graubunden", "GR"), ("grisons", "GR"),
                ("schaffhausen", "SH"), ("glarus", "GL"), ("nidwalden", "NW"), ("obwalden", "OW"), ("uri", "UR"),
                ("appenzell", "AR")]
DEPARTEMENTS = {
    "01": "Ain", "02": "Aisne", "03": "Allier", "04": "Alpes-de-Haute-Provence", "05": "Hautes-Alpes",
    "06": "Alpes-Maritimes", "07": "Ardèche", "08": "Ardennes", "09": "Ariège", "10": "Aube", "11": "Aude",
    "12": "Aveyron", "13": "Bouches-du-Rhône", "14": "Calvados", "15": "Cantal", "16": "Charente",
    "17": "Charente-Maritime", "18": "Cher", "19": "Corrèze", "2A": "Corse-du-Sud", "2B": "Haute-Corse",
    "21": "Côte-d'Or", "22": "Côtes-d'Armor", "23": "Creuse", "24": "Dordogne", "25": "Doubs", "26": "Drôme",
    "27": "Eure", "28": "Eure-et-Loir", "29": "Finistère", "30": "Gard", "31": "Haute-Garonne", "32": "Gers",
    "33": "Gironde", "34": "Hérault", "35": "Ille-et-Vilaine", "36": "Indre", "37": "Indre-et-Loire", "38": "Isère",
    "39": "Jura", "40": "Landes", "41": "Loir-et-Cher", "42": "Loire", "43": "Haute-Loire", "44": "Loire-Atlantique",
    "45": "Loiret", "46": "Lot", "47": "Lot-et-Garonne", "48": "Lozère", "49": "Maine-et-Loire", "50": "Manche",
    "51": "Marne", "52": "Haute-Marne", "53": "Mayenne", "54": "Meurthe-et-Moselle", "55": "Meuse", "56": "Morbihan",
    "57": "Moselle", "58": "Nièvre", "59": "Nord", "60": "Oise", "61": "Orne", "62": "Pas-de-Calais",
    "63": "Puy-de-Dôme", "64": "Pyrénées-Atlantiques", "65": "Hautes-Pyrénées", "66": "Pyrénées-Orientales",
    "67": "Bas-Rhin", "68": "Haut-Rhin", "69": "Rhône", "70": "Haute-Saône", "71": "Saône-et-Loire", "72": "Sarthe",
    "73": "Savoie", "74": "Haute-Savoie", "75": "Paris", "76": "Seine-Maritime", "77": "Seine-et-Marne",
    "78": "Yvelines", "79": "Deux-Sèvres", "80": "Somme", "81": "Tarn", "82": "Tarn-et-Garonne", "83": "Var",
    "84": "Vaucluse", "85": "Vendée", "86": "Vienne", "87": "Haute-Vienne", "88": "Vosges", "89": "Yonne",
    "90": "Territoire de Belfort", "91": "Essonne", "92": "Hauts-de-Seine", "93": "Seine-Saint-Denis",
    "94": "Val-de-Marne", "95": "Val-d'Oise"}
REGIONS_FR = [("bourgogne", "Bourgogne-Franche-Comté"), ("grand est", "Grand Est"), ("brittany", "Bretagne"),
              ("bretagne", "Bretagne"), ("ile de france", "Île-de-France"), ("idf", "Île-de-France"),
              ("occitanie", "Occitanie"), ("languedoc", "Occitanie"), ("hauts de france", "Hauts-de-France"),
              ("auvergne", "Auvergne-Rhône-Alpes"), ("provence", "Provence-Alpes-Côte d'Azur"),
              ("normand", "Normandie"), ("aquitaine", "Nouvelle-Aquitaine"), ("pays de la loire", "Pays de la Loire"),
              ("centre", "Centre-Val de Loire"), ("cors", "Corse")]
LANDER = [("baden w", "Bade-Wurtemberg"), ("bavaria", "Bavière"), ("bayern", "Bavière"), ("berlin", "Berlin"),
          ("brandenburg", "Brandebourg"), ("bremen", "Brême"), ("hamburg", "Hambourg"), ("hesse", "Hesse"),
          ("hessen", "Hesse"), ("lower saxony", "Basse-Saxe"), ("niedersachsen", "Basse-Saxe"),
          ("mecklenburg", "Mecklembourg-Poméranie-Occidentale"), ("north rhine", "Rhénanie-du-Nord-Westphalie"),
          ("nordrhein", "Rhénanie-du-Nord-Westphalie"), ("rhineland", "Rhénanie-Palatinat"),
          ("rheinland", "Rhénanie-Palatinat"), ("saarland", "Sarre"), ("saxony anhalt", "Saxe-Anhalt"),
          ("sachsen anhalt", "Saxe-Anhalt"), ("saxony", "Saxe"), ("sachsen", "Saxe"),
          ("schleswig", "Schleswig-Holstein"), ("thuringia", "Thuringe"), ("thuringen", "Thuringe")]
REGIONS_IT = [("lombard", "Lombardie"), ("marche", "Marches"), ("piedmont", "Piémont"), ("piemonte", "Piémont"),
              ("veneto", "Vénétie"), ("emilia", "Émilie-Romagne"), ("tuscan", "Toscane"), ("toscana", "Toscane"),
              ("lazio", "Latium"), ("liguria", "Ligurie"), ("trentino", "Trentin-Haut-Adige"),
              ("friuli", "Frioul-Vénétie Julienne"), ("campania", "Campanie"), ("puglia", "Pouilles"),
              ("apulia", "Pouilles"), ("sicil", "Sicile"), ("sardegna", "Sardaigne"), ("sardinia", "Sardaigne"),
              ("abruzzo", "Abruzzes"), ("umbria", "Ombrie"), ("calabria", "Calabre"), ("aosta", "Vallée d'Aoste")]
LEGAL = {"sa", "ag", "gmbh", "sarl", "srl", "sas", "sasu", "spa", "s", "a", "r", "l", "co", "kg", "ltd", "inc", "llc",
         "bv", "nv", "ab", "oy", "se", "eurl", "the", "group", "groupe", "holding"}
# registres du commerce publics (domaines) et sites d'emploi courants
REGISTRES = ("zefix", "shab", "fosc", "sogc", "handelsregister", "infogreffe", "pappers", "bodacc", "registroimprese")
EMPLOI = ("jobs.ch", "jobup.ch", "jobscout24", "indeed", "stepstone", "glassdoor", "xing.com/jobs", "monster",
          "apec.fr", "welcometothejungle", "linkedin.com/jobs")
STATUTS_SURS = ("confirmé", "probable")
ORDRE_STATUTS = {"confirmé": 0, "probable": 1, "indice": 2, "aucun": 3}


# ---------------------------------------------------------------- petits outils

def plain(s: str) -> str:
    s = unicodedata.normalize("NFKD", (s or "").lower()).encode("ascii", "ignore").decode()
    return re.sub(r"[^a-z0-9]+", " ", s).strip()


def clip(s: str, n: int) -> str:
    s = re.sub(r"\s+", " ", s or "").strip()
    return s if len(s) <= n else s[: n - 1].rstrip() + "…"


def cle_linkedin(url: str) -> str:
    u = unquote(url or "").lower().strip().rstrip("/")
    u = re.sub(r"^https?://", "", u)
    u = re.sub(r"^[a-z]{2,3}\.linkedin", "linkedin", u)
    return re.sub(r"^www\.", "", u)


def ym(s: str) -> tuple[int, int] | None:
    m = re.match(r"(\d{4})(?:-(\d{1,2}))?", s or "")
    return (int(m[1]), int(m[2] or 1)) if m else None


def fmt_date(s: str) -> str:
    m = re.match(r"(\d{4})(?:-(\d{1,2}))?", s or "")
    if not m:
        return "s.d."
    return f"{int(m[2]):02d}.{m[1]}" if m[2] else m[1]


def mois_depuis(s: str) -> int | None:
    d = ym(s)
    return None if not d else (REF.year - d[0]) * 12 + REF.month - d[1]


def duree(mois: int) -> str:
    if mois < 12:
        return f"{max(mois, 1)} mois"
    a, m = divmod(mois, 12)
    return f"{a} an{'s' if a > 1 else ''}" + (f" {m} mois" if m else "")


def millions(v: str) -> str:
    x = float(v) / 1e6
    return (f"{x:.0f}" if x >= 10 else f"{x:.1f}".replace(".", ",")) + " M"


def nombre(v: str) -> float:
    try:
        return float(v or 0)
    except ValueError:
        return 0.0


def mots_nom(nom: str) -> list[str]:
    """Mots distinctifs d'un nom d'entreprise : 'Exemple Precision SA' -> ['exemple', 'precision']."""
    s = re.sub(r"\(.*?\)|\(.*$", "", nom or "")
    return [t for t in plain(s).split() if t not in LEGAL and len(t) >= 3]


def fmt_tel(t: str) -> str:
    """'+41270000000' -> '+41 27 000 00 00' ; '+33600000000' -> '+33 6 00 00 00 00' ; le reste tel quel."""
    n = re.sub(r"[^\d+]", "", t or "")
    if re.fullmatch(r"\+41\d{9}", n):
        return f"{n[:3]} {n[3:5]} {n[5:8]} {n[8:10]} {n[10:]}"
    if re.fullmatch(r"\+33\d{9}", n):
        return f"{n[:3]} {n[3]} " + " ".join(n[i:i + 2] for i in range(4, 12, 2))
    return (t or "").strip()


def sans_tirets(s: str) -> str:
    return re.sub(r"\s*[\u2014\u2013]\s*", ", ", s or "")


def date_fichier(p: Path) -> str:
    return datetime.fromtimestamp(p.stat().st_mtime).strftime("%d.%m.%Y") if p.exists() else "absent"


# ---------------------------------------------------------------- données

class Donnees:
    def __init__(self, camp: Campagne, a: argparse.Namespace) -> None:
        global REF
        self.camp, self.cfg = camp, camp.config
        REF = camp.date_reference()
        self.langues = [l.strip().lower() for l in camp.config["appels"].get("langues") or [] if l.strip()]
        self.chemin_ordre = Path(a.ordre) if a.ordre else camp.appels / "ordre-appel.csv"
        if not self.chemin_ordre.exists():
            raise SystemExit(f"{self.chemin_ordre} absent : lancer d'abord ordre_appel.py (étape 8).")
        self.entree = [self._contact(r) for r in lire_csv(self.chemin_ordre)]
        self.fiches = charger_fiches(camp, obligatoire=True)
        self.entreprises, manquantes = {}, []
        for c in self.entree:
            e = self.fiches.de(c)
            if not e:
                manquantes.append(c["entreprise"])
                continue
            c["_id"] = e["id"]
            self.entreprises[e["id"]] = self._fiche(e)
        if manquantes:
            print(f"! {len(manquantes)} contact(s) écarté(s), entreprise absente de prospection-signaux.csv : "
                  f"{sorted(set(manquantes))} (corriger entreprise_id ou le nom dans {self.chemin_ordre.name})")
            self.entree = [c for c in self.entree if c.get("_id")]
        if not self.entree:
            raise SystemExit(f"Aucun contact de {self.chemin_ordre.name} ne se rattache à prospection-signaux.csv.")
        sig = camp.dossier / "donnees" / "signaux"
        self.signaux = lire_json(sig / "signaux.json", {}) or {}
        self.recherche = lire_json(sig / "recherche.json", {}) or {}
        self.linkedin = lire_json(sig / "linkedin.json", {}) or {}
        self.logiciels = lire_json(sig / "logiciels.json", {}) or {}
        self.entites = {e.get("id"): e for e in lire_json(sig / "entites.json", []) or [] if isinstance(e, dict)}
        self.ui = {}
        if a.export:
            self.ui = {r.get("Apollo Contact Id", ""): r for r in lire_csv(a.export) if r.get("Apollo Contact Id")}
        self.comptes, self.comptes_nom = comptes_apollo(camp)
        # relais, groupés par entreprise
        self.relais = defaultdict(list)
        try:
            relais, _ = charger_contacts(camp, a.relais, ORDRE_RELAIS)
        except SystemExit as e:
            print(f"! relais : {e}")
            relais = []
        for r in relais:
            e = self.fiches.de(r)
            if e and e["id"] not in self.entreprises:
                self.entreprises[e["id"]] = self._fiche(e)
            self.relais[e["id"] if e else "nom:" + norm(r["entreprise"])].append(r)
        self.cartes = charger_cartes(camp)
        # prénoms connus et noms de famille par entreprise (repérer un prénom et un nom inversés par Apollo)
        try:
            tous, _ = charger_contacts(camp)
        except SystemExit:
            tous = []
        self.prenoms = defaultdict(int)
        self.noms_entreprise = defaultdict(list)
        for r in tous or self.entree + relais:
            if plain(r["prenom"]):
                self.prenoms[plain(r["prenom"]).split()[0]] += 1
            if plain(r["nom"]):
                cle = (self.fiches.de(r) or {}).get("id") or "nom:" + norm(r["entreprise"])
                self.noms_entreprise[cle].append(plain(r["nom"]).split()[0])
        self.corrections = lire_json(camp.appels / "corrections.json", {}) or {}
        suisses = any(e["pays"] == "CH" for e in self.entreprises.values())
        self.loc = charger_localites() if suisses else {}
        self.par_li, self.par_nom = {}, defaultdict(list)
        cache = camp.dossier / "apollo" / "cache"
        for f in sorted(cache.glob("*.json")) if cache.is_dir() else []:
            try:
                d = json.loads(f.read_text(encoding="utf-8"))
            except (OSError, json.JSONDecodeError):
                continue
            if isinstance(d, dict) and isinstance(d.get("matches"), list):
                for p in d["matches"]:
                    if not p:
                        continue
                    if p.get("linkedin_url"):
                        self.par_li[cle_linkedin(p["linkedin_url"])] = p
                    self.par_nom[(plain(p.get("first_name")), plain(p.get("last_name")))].append(p)

    @staticmethod
    def _contact(r: dict) -> dict:
        """Ligne de l'ordre d'appel ; contact_id = clé du contact (id Apollo ou id local), id_apollo = id Apollo seul."""
        c = dict(r)
        c["contact_id"] = cle_contact(r)
        c["id_apollo"] = id_apollo(c["contact_id"])
        for k in ("prenom", "nom", "titre", "email", "linkedin", "entreprise", "entreprise_id", "deja_contacte",
                  "a_eviter", "telephone", "type_telephone", "langue", "parle_langue_appel", "langue_principale",
                  "fenetre_opportunite", "cibles", "priorite", "ville", "pays"):
            c[k] = r.get(k) or ""
        return c

    @staticmethod
    def _fiche(e: dict) -> dict:
        f = {k: e.get(k) or "" for k in ("id", "nom", "tier", "canton", "secteur", "taille", "activite",
                                         "pourquoi_maintenant", "angle", "dirigeants", "signaux_tech", "linkedin")}
        f["pays"] = pays_entreprise(e)
        f["site"] = e.get("site") or e.get("website") or ""
        return f

    def personne(self, c: dict) -> dict | None:
        """Fiche Apollo enrichie de la personne (LinkedIn d'abord, puis prénom et nom, dans les deux sens)."""
        p = self.par_li.get(cle_linkedin(c.get("linkedin", "")))
        if p:
            return p
        for a, b in ((c["prenom"], c["nom"]), (c["nom"], c["prenom"])):
            cands = self.par_nom.get((plain(a), plain(b)), [])
            if len(cands) == 1:
                return cands[0]
        return None

    def canton_ville(self, ville: str) -> str:
        """Canton d'une localité suisse (répertoire des localités), '' si inconnue."""
        k = norm(ville)
        v = self.loc.get(k) or self.loc.get(re.sub(r"\b(st|ste)\b", "saint", k))
        return v[0] if v else ""

    def ui_entreprise(self, eid: str) -> dict:
        """Ligne de l'export de l'interface Apollo d'un contact de l'entreprise (champs « Company ... »)."""
        for c in self.entree:
            if c.get("_id") == eid and c["contact_id"] in self.ui:
                return self.ui[c["contact_id"]]
        return {}

    def compte(self, e: dict) -> dict:
        """Compte Apollo de l'entreprise : par domaine, sinon par nom."""
        return (self.comptes.get(domaine(e["site"])) or self.comptes_nom.get(norm(e["nom"]))
                or self.comptes_nom.get(norm(self.ui_entreprise(e["id"]).get("Company Name", ""))) or {})

    def logiciel(self, eid: str) -> dict:
        """Logiciel relevé pour l'entreprise (logiciels.json, relu à la main) ; le plus sûr s'il y en a plusieurs."""
        v = self.logiciels.get(eid)
        if isinstance(v, list):
            v = min((x for x in v if isinstance(x, dict)), key=lambda x: ORDRE_STATUTS.get(x.get("statut", ""), 9), default=None)
        return v if isinstance(v, dict) else {}

    def nom_inverse(self, c: dict, cle: str) -> bool:
        """Apollo inverse parfois prénom et nom. Inversé si le « nom » est un prénom connu et pas le « prénom », ET
        si l'email (p.nom@) ou un collègue du même nom de famille le confirme."""
        p, n = (plain(c["prenom"]).split() or [""])[0], (plain(c["nom"]).split() or [""])[0]
        if not p or not n or not (self.prenoms[n] - (n == p) >= 1 and self.prenoms[p] <= 1):
            return False
        local = plain((c.get("email") or "").split("@")[0]).split()
        if local:
            return len(local) >= 2 and local[-1] == p and local[0] in (n, n[0])
        return self.noms_entreprise[cle].count(p) >= 1

    def corriger_nom(self, c: dict, e: dict | None) -> tuple[dict, str]:
        """Prénom et nom tels qu'il faut les dire : inversion d'Apollo corrigée ; nom tronqué (« N ») complété par
        l'email (prenom.nom@) ou, à défaut, par le registre (même prénom, même initiale)."""
        c = dict(c)
        if self.nom_inverse(c, (e or {}).get("id") or "nom:" + norm(c.get("entreprise", ""))):
            c["prenom"], c["nom"] = c["nom"], c["prenom"]
            return c, "Prénom et nom inversés dans Apollo : corrigés ici."
        initiale = c["nom"].strip().rstrip(".")
        prenom = (plain(c["prenom"]).split() or [""])[0]
        if len(initiale) == 1 and prenom:
            local = plain((c.get("email") or "").split("@")[0]).split()
            if len(local) == 2 and local[0] == prenom and len(local[1]) > 1 and local[1][0] == initiale.lower():
                c["nom"] = local[1].capitalize()
                return c, f"Nom tronqué dans Apollo (« {initiale} ») : complété d'après l'email."
            for r in roles_registre(self, e) if e else []:
                mots = (r.get("nom") or "").split()
                if len(mots) >= 2 and plain(mots[0]) == prenom and plain(mots[-1])[:1] == initiale.lower():
                    c["nom"] = " ".join(mots[1:])
                    return c, f"Nom tronqué dans Apollo (« {initiale} ») : complété d'après le registre ({r.get('role', '')})."
        return c, ""


# ---------------------------------------------------------------- lieu

def canton_de(texte: str) -> str:
    t = plain(texte)
    if t.upper() in CANTONS:
        return t.upper()
    return next((code for alias, code in CANTON_ALIAS if alias in t), "")


def chercher(table: list[tuple[str, str]], texte: str) -> str:
    t = plain(texte)
    return next((nom for alias, nom in table if alias in t), "") if t else ""


def localisation(d: Donnees, e: dict) -> dict:
    """Pays, canton ou département, ville de l'entreprise. Apollo et LinkedIn ne servent que s'ils sont dans le
    même pays que la fiche de l'étape 5 (sinon c'est souvent le siège du groupe)."""
    ui = d.ui_entreprise(e["id"])
    li = d.linkedin.get(e["id"], {}) or {}
    compte = d.compte(e)
    code = e["pays"] or code_pays(ui.get("Company Country", "")) or code_pays(li.get("pays", ""))
    a_ok = code_pays(ui.get("Company Country") or compte.get("country") or "") == code
    l_ok = code_pays(li.get("pays", "")) == code
    etat = (ui.get("Company State") or compte.get("state") or "") if a_ok else ""
    li_region = li.get("region", "") if l_ok else ""
    ville = ((li.get("ville") if l_ok else "") or ((ui.get("Company City") or compte.get("city")) if a_ok else "")
             or (d.entites.get(e["id"]) or {}).get("ville") or "")
    ville = re.sub(r"^\d{4,5}\s+", "", ville.strip())   # « 1950 Sion » -> « Sion »
    region, canton = "", ""
    if code == "CH":
        c = e["canton"] or canton_de(li_region) or canton_de(etat)
        cv = d.canton_ville(ville) if ville else ""
        canton = cv or c
        region = f"{CANTONS[canton]} ({canton})" if canton in CANTONS else ""
        if cv and c in CANTONS and cv != c:     # le siège au registre n'est pas sur le site (holding, transfert)
            region += f" ; siège au registre : {CANTONS[c]} ({c})"
    elif code == "FR":
        cp = (compte.get("postal_code") or "") if a_ok else ""
        dep = "2A" if cp.startswith("200") else "2B" if cp.startswith("202") else cp[:2]
        region = f"{DEPARTEMENTS[dep]} ({dep})" if dep in DEPARTEMENTS else chercher(REGIONS_FR, f"{li_region} {etat}")
        if region and dep not in DEPARTEMENTS:
            region += " (région)"
    elif code == "DE":                         # LinkedIn d'abord : Apollo confond les villes homonymes
        region = chercher(LANDER, li_region) or chercher(LANDER, etat)
    elif code == "IT":
        region = chercher(REGIONS_IT, li_region) or chercher(REGIONS_IT, etat)
    note = ""
    if not a_ok and ui.get("Company Country"):
        note = f"compte Apollo rattaché à {ui.get('Company City') or '?'} ({ui.get('Company Country')})"
    return {"pays": NOMS_PAYS.get(code, code or "inconnu"), "code": code, "region": region or "inconnu",
            "ville": ville.strip(), "note": note, "canton": canton}


# ---------------------------------------------------------------- personne

def meme_entreprise(org: str, e: dict) -> bool:
    mots = set(mots_nom(e["nom"])) | set(mots_nom(re.sub(r"^https?://|www\.|\.[a-z]+$", "", domaine(e["site"]))))
    return bool(mots & set(plain(org).split()))


def roles_registre(d: Donnees, e: dict) -> list[dict]:
    rech = d.recherche.get(e["id"], {}) or {}
    if rech.get("dirigeants"):
        return rech["dirigeants"]
    out, prof, cur = [], 0, ""
    for ch in e["dirigeants"] + ";":          # « Nom (rôle ; détail, depuis 2025-02); Nom2 (...) »
        if ch == "(":
            prof += 1
        elif ch == ")":
            prof -= 1
        if ch == ";" and prof == 0:
            m = re.match(r"\s*([^(]+?)\s*\((.*)\)\s*$", cur)
            if m:
                role, depuis = m[2], ""
                md = re.search(r",\s*depuis (\d{4}(?:-\d{2})?)\s*$", role)
                if md:
                    role, depuis = role[: md.start()], md[1]
                out.append({"nom": m[1], "role": role, "depuis": depuis})
            cur = ""
        else:
            cur += ch
    return out


def cite(nom_personne: str, prenom: str, texte: str) -> bool:
    """Le texte nomme-t-il la personne ? Nom ET prénom exigés : un même nom de famille désigne souvent le père
    ou le frère dans une entreprise familiale."""
    t = set(plain(texte).split())
    noms = [x for x in plain(nom_personne).split() if len(x) >= 3]
    prenoms = plain(prenom).split()
    return bool(noms) and all(x in t for x in noms) and (not prenoms or prenoms[0] in t)


def proximite(titre_a: str, titre_b: str) -> int:
    return len(set(plain(titre_a).split()) & set(plain(titre_b).split()))


def annees(h: dict) -> str:
    a, b = (h.get("start_date") or "")[:4], (h.get("end_date") or "")[:4]
    return f" ({a})" if a and a == b else f" ({a}-{b})" if a and b else f" (depuis {a})" if a else ""


def role_officiel(d: Donnees, c: dict, e: dict) -> dict | None:
    """Rôle de la personne au registre ou cité publiquement (fiche de l'étape 5)."""
    for r in roles_registre(d, e):
        if cite(r.get("nom", ""), c["prenom"], f"{c['prenom']} {c['nom']}") or cite(c["nom"], c["prenom"], r.get("nom", "")):
            return r
    return None


def faits_personne(d: Donnees, c: dict, e: dict, loc: dict) -> list[str]:
    """Faits sur la personne : poste et ancienneté, parcours (Apollo, d'après LinkedIn), rôle au registre ou cité
    publiquement, signal qui la nomme, lieu de vie s'il diffère, historique avec nous."""
    faits = []
    p = d.personne(c)
    if p:
        hist = [h for h in (p.get("employment_history") or []) if h.get("title") or h.get("organization_name")]
        ici = [h for h in hist if (p.get("organization_id") and h.get("organization_id") == p.get("organization_id"))
               or meme_entreprise(h.get("organization_name", ""), e)]
        actuels = [h for h in ici if h.get("current")]
        # le poste actuel est celui dont l'intitulé ressemble le plus à celui de la liste, puis le plus récent
        actuels.sort(key=lambda h: (proximite(h.get("title", ""), c["titre"]), h.get("start_date") or ""), reverse=True)
        principal = actuels[0] if actuels else None
        if principal and principal.get("start_date"):
            m = mois_depuis(principal["start_date"])
            debut = "Nouveau à ce poste" if m is not None and m <= NOUVEAU_POSTE_MOIS else "En poste"
            faits.append(f"{debut} : {principal['title']} depuis {fmt_date(principal['start_date'])} ({duree(m or 0)}).")
        elif principal or ici:
            faits.append(f"Poste : {(principal or ici[0]).get('title') or c['titre']} (date de prise de poste inconnue).")
        for h in actuels[1:2]:
            faits.append(f"Aussi : {h['title']}{annees(h)}.")
        anciens_ici = sorted([h for h in ici if not h.get("current") and h.get("start_date")], key=lambda h: h["start_date"])
        if anciens_ici and principal and principal.get("start_date"):
            premier = anciens_ici[0]
            fin = premier.get("end_date") or ""
            ecart = (mois_depuis(fin) or 0) - (mois_depuis(principal["start_date"]) or 0) if fin else 0
            if ecart >= 12:              # parti puis revenu : ce n'est pas une carrière continue dans la maison
                faits.append(f"Déjà passé par l'entreprise : {premier['title']}{annees(premier)}.")
            elif premier["start_date"] < principal["start_date"]:
                faits.append(f"Dans l'entreprise depuis {fmt_date(premier['start_date'])} (entré comme {premier['title']}).")
        for h in [h for h in hist if h.get("current") and h not in ici][:2]:
            faits.append(f"Autre poste en cours : {h['title']} chez {h['organization_name']}{annees(h)}.")
        passes = sorted([h for h in hist if not h.get("current") and h not in ici and h.get("organization_name")],
                        key=lambda h: h.get("end_date") or h.get("start_date") or "", reverse=True)[:2]
        if passes:
            faits.append("Avant : " + " ; ".join(f"{h['title']} chez {h['organization_name']}{annees(h)}"
                                                 for h in passes) + ".")
    r = role_officiel(d, c, e)
    if r:
        depuis = f" (depuis {fmt_date(r['depuis'])})" if r.get("depuis") else ""
        faits.append(f"Rôle officiel ou public : {sans_tirets(r.get('role', '').strip())}{depuis}.")
    nommes = sorted([s for s in d.signaux.get(e["id"], []) or []
                     if s.get("famille") != "negatif" and cite(c["nom"], c["prenom"], s.get("resume", ""))],
                    key=lambda s: s.get("date") or "", reverse=True)
    if nommes:
        faits.append(f"Cité dans un signal : {sans_tirets(nommes[0]['resume'])} ({fmt_date(nommes[0].get('date'))}).")
    ui = d.ui.get(c.get("contact_id", ""), {})
    ville_p = ui.get("City") or c.get("ville") or (p or {}).get("city") or ""
    pays_p = code_pays(ui.get("Country") or c.get("pays") or (p or {}).get("country") or "")
    if ville_p and loc:
        ailleurs = (pays_p and pays_p != loc["code"]) or (
            loc["code"] == "CH" and d.canton_ville(ville_p) and d.canton_ville(ville_p) != loc.get("canton"))
        if ailleurs:
            faits.append(f"Basé à {ville_p}" + (f" ({NOMS_PAYS.get(pays_p, pays_p)})" if pays_p != loc["code"] else "") + ".")
    if c.get("deja_contacte"):
        societe = (d.cfg["offre"].get("societe") or "").strip()
        faits.append(f"Déjà contacté{' par ' + societe if societe else ''} : {c['deja_contacte']}.")
    return [sans_tirets(f) for f in faits]


# ---------------------------------------------------------------- entreprise

def source(s: dict) -> str:
    dom = domaine(s.get("url") or "")
    if any(x in dom for x in REGISTRES):
        lab = "registre"
    elif "linkedin" in dom:
        lab = "LinkedIn"
    else:
        lab = dom or s.get("source", "")
    return lab + (", indice" if s.get("confiance") == "auto" else "")


def ligne_signal(s: dict) -> str:
    return f"{fmt_date(s.get('date'))} : {sans_tirets(s.get('resume', ''))} ({source(s)})"


def signaux(d: Donnees, e: dict, familles: tuple[str, ...], n: int = 4) -> list[str]:
    lst = [s for s in d.signaux.get(e["id"], []) or [] if s.get("famille") in familles]
    lst = sorted(lst, key=lambda s: -nombre(s.get("poids")))[:n]
    lst.sort(key=lambda s: s.get("date") or "", reverse=True)
    return [ligne_signal(s) for s in lst]


def direction(d: Donnees, e: dict) -> list[str]:
    out = signaux(d, e, ("direction", "ma"))
    parent = (d.recherche.get(e["id"], {}) or {}).get("groupe_parent") or ""
    if not parent:
        apollo_parent = d.ui_entreprise(e["id"]).get("Parent company (Apollo data)", "")
        parent = f"{apollo_parent} (selon Apollo, à vérifier)" if apollo_parent else ""
    if parent:
        out.insert(0, f"Groupe ou actionnaire : {sans_tirets(parent)}")
    return out


def finances(d: Donnees, e: dict) -> list[str]:
    out = signaux(d, e, ("financement",))
    ui = d.ui_entreprise(e["id"])
    if nombre(ui.get("Annual Revenue")) > 0:
        out.append(f"Chiffre d'affaires estimé par Apollo : {millions(ui['Annual Revenue'])} USD (estimation, non vérifiée)")
    if ui.get("Latest Funding"):
        txt = f"Apollo signale une dernière opération « {ui['Latest Funding']} »"
        if nombre(ui.get("Latest Funding Amount")) > 0:
            txt += f" de {millions(ui['Latest Funding Amount'])} USD"
        if ui.get("Last Raised At"):
            txt += f" ({fmt_date(ui['Last Raised At'][:7])})"
        if nombre(ui.get("Total Funding")) > 0:
            txt += f", {millions(ui['Total Funding'])} USD levés au total"
        out.append(txt + " (à vérifier)")
    return out


def tech(d: Donnees, e: dict) -> list[str]:
    out = []
    at = (d.recherche.get(e["id"], {}) or {}).get("adoption_tech")
    if at:
        out.append(f"Adoption tech : {at.get('niveau', '?')}" + (" ; vend une offre concurrente de la nôtre" if at.get("concurrent") else ""))
        out += [sans_tirets(x) for x in at.get("preuves", [])[:3]]
    else:
        out += [x.strip() for x in e["signaux_tech"].split(" | ") if x.strip()][:4]
    li = d.linkedin.get(e["id"], {}) or {}
    if at and li.get("posts_90j") not in (None, "", "0", 0):
        out.append(f"LinkedIn : {li['posts_90j']} posts en 90 jours")
    techs = [t.strip() for t in d.ui_entreprise(e["id"]).get("Technologies", "").split(",") if t.strip()]
    vus = [t for nom in d.cfg["signaux"].get("logiciels") or [] for t in techs
           if re.match(rf"{re.escape(nom)}\b", t, re.I)]
    vus = list(dict.fromkeys(vus))
    if vus:
        out.append("Logiciels repérés par Apollo (indices) : " + ", ".join(vus[:10]))
    return out


def logiciels_txt(d: Donnees, e: dict) -> str:
    v = d.logiciel(e["id"])
    projets = [s for s in d.signaux.get(e["id"], []) or [] if s.get("sous_type") == "erp_numerique"]
    projet = ""
    if projets:
        s = max(projets, key=lambda s: s.get("date") or "")
        projet = f"\nProjet numérique signalé : {sans_tirets(s.get('resume', ''))} ({fmt_date(s.get('date'))})."
    if not v.get("logiciel") or v.get("statut") == "aucun":
        return "Inconnu : à demander au téléphone." + projet
    src = sans_tirets(v.get("source") or domaine(v.get("url") or "") or "source à relire")
    return f"{v['logiciel']} ({v.get('statut') or 'statut inconnu'}) : {src}." + projet


def pourquoi_maintenant(d: Donnees, e: dict, redige: str = "") -> str:
    """Texte de l'étape 5, corrigé s'il citait un signal négatif (appels/corrections.json). Le repli « Tech : ... »
    n'est pas un déclencheur : remplacé par le dernier fait daté de la fenêtre (signaux.fenetre_mois), s'il existe."""
    txt = ((d.corrections.get(e["id"], {}) or {}).get("pourquoi_maintenant") or e["pourquoi_maintenant"] or redige or "").strip()
    if txt and not txt.startswith("Tech"):
        return sans_tirets(txt)
    fenetre = int(d.cfg["signaux"].get("fenetre_mois") or 24)
    recents = [s for s in d.signaux.get(e["id"], []) or [] if s.get("famille") != "negatif"
               and (mois_depuis(s.get("date") or "") or 999) <= fenetre]
    txt = "Pas de déclencheur daté : appel de découverte."
    if recents:
        s = max(recents, key=lambda s: (s.get("date") or "", nombre(s.get("poids"))))
        txt += f" Dernier fait daté : {sans_tirets(s.get('resume', ''))} ({fmt_date(s.get('date'))})."
    return txt


def taille_sources(d: Donnees, e: dict) -> str:
    li = d.linkedin.get(e["id"], {}) or {}
    ui = d.ui_entreprise(e["id"])
    parts = [f"Étape 5 : {e['taille'] or 'inconnue'}"]
    if li.get("taille"):
        parts.append(f"LinkedIn : {li['taille']}, {li.get('salaries_linkedin') or '?'} salariés inscrits")
    if ui.get("# Employees"):
        parts.append(f"Apollo : {ui['# Employees']} employés")
    return " | ".join(parts)


CODES_LANGUES = {plain(v): k for k, v in NOMS_LANGUES.items()}


def langue_txt(c: dict, langues: list[str]) -> str:
    """Colonne Langue : pour « oui » et « probable », une langue de l'appelant (celle de langue_principale si
    l'appelant la parle, sinon sa première langue) ; pour « incertain » et « autre langue », la langue des indices."""
    classe, brut = c.get("parle_langue_appel", ""), (c.get("langue_principale") or "").strip()
    code = brut.lower() if brut.lower() in NOMS_LANGUES else CODES_LANGUES.get(plain(brut), "")
    if classe in ("oui", "probable") and langues and code not in langues:
        code, brut = langues[0], ""
    lp = NOMS_LANGUES.get(code) or brut or "?"
    return {"oui": lp, "probable": f"{lp} (probable)", "incertain": f"incertain ({lp} ?)",
            "non": lp if lp != "?" else "autre langue"}.get(classe, c.get("langue") or lp)


def telephone_txt(c: dict) -> str:
    if not c.get("telephone"):
        return "aucun"
    return f"{fmt_tel(c['telephone'])} ({c.get('type_telephone') or '?'})"


def standard(d: Donnees, e: dict, code: str) -> str:
    ui = d.ui_entreprise(e["id"])
    tel = ui.get("Company Phone") or d.compte(e).get("sanitized_phone") or ""
    tel = re.sub(r"^'", "", tel.strip())
    if tel and code in INDICATIF and not re.sub(r"[^\d+]", "", tel).startswith(INDICATIF[code]):
        return ""                              # standard de groupe ou de filiale à l'étranger : inutile ici
    return fmt_tel(tel)


def cibles_txt(r: dict) -> str:
    """'Prénom Nom (qualité et méthodes)' -> 'qualité et méthodes'."""
    t = r.get("cibles") or ""
    m = re.fullmatch(r"\s*([^()]+?)\s*\((.*)\)\s*", t)
    if m and set(plain(m[1]).split()) & set(plain(f"{r['prenom']} {r['nom']}").split()):
        return m[2]
    return t


FAMILLES = {"direction": "direction", "qualite": "qualité", "it": "informatique", "production": "production",
            "technique": "bureau technique", "service": "service client", "ventes": "ventes", "finance": "finances",
            "achats": "achats", "projet": "projets"}


def relais_par_defaut(cfg: dict) -> str:
    familles = cfg["persona"].get("relais") or []
    familles = [familles] if isinstance(familles, str) else familles
    noms = [FAMILLES.get(f, f) for f in familles]
    return "aucun relais trouvé : demander qui gère au quotidien" + (" : " + ", ".join(noms) if noms else " ce sujet")


def relais_txt(d: Donnees, eid: str) -> str:
    lignes = []
    for r in d.relais.get(eid, []):
        r, _ = d.corriger_nom(r, d.entreprises.get(eid))
        bits = [f"{r['prenom']} {r['nom']}, {r['titre']}"]
        if cibles_txt(r):
            bits[0] += f" ({cibles_txt(r)})"
        if r.get("email"):
            bits.append(r["email"])
        if r.get("mobile") or r.get("telephone_standard"):
            bits.append(fmt_tel(r.get("mobile")) or f"standard {fmt_tel(r['telephone_standard'])}")
        lignes.append(" ; ".join(bits))
    return "\n".join(lignes) or relais_par_defaut(d.cfg)


# ---------------------------------------------------------------- briefs pour les agents

# lots écrits par lot d'agent (L<n>.json) : un nouveau découpage les rend faux, un contact pouvant changer de lot
LOTS_AGENTS = ("briefs", "redaction", "cartes", "sources/briefs", "sources/redaction")


def cle_json(x: dict) -> str:
    """Clé d'un contact dans un brief ou une réponse d'agent (contact_id recopié) ; '' si absente."""
    return (x.get("contact_id") or x.get("apollo_contact_id") or "").strip()


def composition(lot: list) -> list:
    """Entreprises et contacts d'un lot de briefs, pour comparer deux découpages."""
    return [(x.get("id"), [cle_json(c) for c in x.get("contacts") or []]) for x in lot if isinstance(x, dict)]


def composition_lots(dossier: Path) -> dict:
    return {f.name: composition(lire_json(f, []) or []) for f in sorted(dossier.glob("L*.json"))}


def archiver_lots(appels: Path) -> None:
    """Range les lots de la passe précédente (briefs, rédactions, cartes, sources) dans appels/precedent/<date-heure>/."""
    fichiers = [f for sous in LOTS_AGENTS for f in sorted((appels / sous).glob("L*.json"))]
    if not fichiers:
        return
    dest = appels / "precedent" / datetime.now().strftime("%Y%m%d-%H%M%S")
    for f in fichiers:
        cible = dest / f.relative_to(appels)
        cible.parent.mkdir(parents=True, exist_ok=True)
        f.replace(cible)
    agents = sorted({str(f.parent.relative_to(appels)) for f in fichiers if f.parent.name != "briefs"})
    print(f"! découpage en lots changé : passe précédente rangée dans {dest} ({len(fichiers)} fichiers)"
          + (f" ; à refaire par les agents sur les nouveaux lots : {', '.join(agents)}" if agents else ""))


def briefs(d: Donnees, n_lots: int) -> None:
    ordre = list(dict.fromkeys(c["_id"] for c in d.entree))
    fiches = []
    for eid in ordre:
        e = d.entreprises[eid]
        loc = localisation(d, e)
        li = d.linkedin.get(eid, {}) or {}
        ui = d.ui_entreprise(eid)
        rech = d.recherche.get(eid, {}) or {}
        contacts = []
        for c in d.entree:
            if c["_id"] != eid:
                continue
            c2, note_nom = d.corriger_nom(c, e)
            carte = d.cartes.get(c["contact_id"])
            contacts.append({
                "contact_id": c["contact_id"], "prenom": c2["prenom"], "nom": c2["nom"], "note_nom": note_nom,
                "titre": c["titre"], "langue": langue_txt(c, d.langues), "deja_contacte": c.get("deja_contacte", ""),
                "faits_personne": faits_personne(d, c2, e, loc),
                "carte_existante": {k: carte[k] for k in ("pourquoi_vous", "exemple_pour_eux", "relance",
                                                          "objection_probable", "preuve") if carte.get(k)} if carte else None,
                **({"absent_d_apollo": True} if not c["id_apollo"] else {}),
            })
        tous = sorted(d.signaux.get(eid, []) or [], key=lambda s: s.get("date") or "", reverse=True)
        fiches.append({
            "id": eid, "entreprise": e["nom"], "tier": e["tier"], "pays": loc["pays"], "code_pays": loc["code"],
            "canton_ou_departement": loc["region"], "ville": loc["ville"],
            "lieux_a_eviter": lieux_a_eviter(loc["code"], d.cfg),
            "categorie": e["secteur"], "activite": rech.get("activite") or e["activite"],
            "linkedin": {"secteur": li.get("secteur", ""), "specialites": clip(li.get("specialites", ""), 300),
                         "description": clip(li.get("description", ""), 600), "fondation": li.get("fondation", "")},
            "apollo": {"industrie": ui.get("Industry", ""), "mots_cles": ", ".join(ui.get("Keywords", "").split(", ")[:15])},
            "effectif_sources": taille_sources(d, e),
            "groupe_parent": rech.get("groupe_parent", ""),
            "signaux_dates": [ligne_signal(s) for s in tous if s.get("famille") != "negatif"][:12],
            "tech": tech(d, e), "finances": finances(d, e), "logiciels": logiciels_txt(d, e),
            "pourquoi_maintenant": e["pourquoi_maintenant"], "angle": e["angle"],
            "a_ne_jamais_evoquer": [ligne_signal(s) for s in tous if s.get("famille") == "negatif"],
            "relais": [f"{r['prenom']} {r['nom']}, {r['titre']} ({cibles_txt(r)})"
                       for r in (d.corriger_nom(x, e)[0] for x in d.relais.get(eid, []))],
            "contacts": contacts,
        })
    # lots équilibrés en nombre de contacts, une entreprise entière par lot
    lots = [[] for _ in range(max(1, min(n_lots, len(fiches))))]
    for f in sorted(fiches, key=lambda f: -len(f["contacts"])):
        min(lots, key=lambda l: sum(len(x["contacts"]) for x in l)).append(f)
    dossier = d.camp.appels / "briefs"
    dossier.mkdir(parents=True, exist_ok=True)
    (d.camp.appels / "redaction").mkdir(parents=True, exist_ok=True)
    for lot in lots:
        lot.sort(key=lambda f: ordre.index(f["id"]))
    nouveau = {f"L{i}.json": composition(lot) for i, lot in enumerate(lots, 1)}
    if composition_lots(dossier) != nouveau:
        archiver_lots(d.camp.appels)
    for ancien in dossier.glob("L*.json"):
        ancien.unlink()
    for i, lot in enumerate(lots, 1):
        (dossier / f"L{i}.json").write_text(json.dumps(lot, ensure_ascii=False, indent=1), encoding="utf-8")
        print(f"L{i} : {len(lot)} entreprises, {sum(len(f['contacts']) for f in lot)} contacts")
    o = d.cfg["offre"]
    langues = [l for l in d.cfg["appels"].get("langues") or [] if l]
    contexte = {
        "campagne": d.camp.nom, "date_reference": REF.isoformat(),
        "offre": {k: o.get(k) for k in ("societe", "description", "argument_cle", "preuves", "appelant",
                                        "ne_jamais_dire", "concurrents")},
        "persona": {"entree": d.cfg["persona"].get("entree"), "relais": d.cfg["persona"].get("relais")},
        "langues_appelant": langues, "langue_redaction": langues[0] if langues else "fr",
        "consignes": {"redaction": "consignes/fichier-appel-redaction.md", "cartes": "consignes/cartes-appel.md",
                      "sources": "consignes/fichier-appel-sources.md"},
    }
    (dossier / "contexte.json").write_text(json.dumps(contexte, ensure_ascii=False, indent=1), encoding="utf-8")
    print(f"-> {dossier} (L1 à L{len(lots)}, contexte.json) ; réponses attendues dans {d.camp.appels / 'redaction'}")


# ---------------------------------------------------------------- sources des arguments (« J'ai vu sur ... que ... »)

def libelle_source(url: str, site: str) -> str:
    """Où l'appelant dira l'avoir vu : registre, LinkedIn, site, offre d'emploi ou article (domaine)."""
    dom, site_dom = domaine(url), domaine(site)
    if not dom:
        return ""
    if any(x in dom for x in REGISTRES):
        return "registre du commerce"
    if "linkedin.com" in dom and "/jobs" not in url:
        return "profil LinkedIn" if "/in/" in url else "page LinkedIn de l'entreprise"
    if any(x in url for x in EMPLOI) or re.search(r"/(jobs?|karriere|stellen|emplois?|careers?|offres?-d-emploi)/", url):
        return f"offre d'emploi ({dom})"
    if site_dom and (dom == site_dom or dom.endswith("." + site_dom) or site_dom.endswith("." + dom)):
        return "site de l'entreprise"
    return f"article ou page ({dom})"


def lien_preuve(preuve: str, sources: list[str]) -> str:
    """Lien d'une preuve tech de l'étape 5 (« ... (post LinkedIn du 2026-09-16) ») parmi les sources de la fiche :
    mots du fait présents dans l'URL, domaine ou LinkedIn cité dans le texte. '' si rien ne se recoupe."""
    mots = {m for m in plain(preuve).split() if len(m) >= 4}
    meilleur, score_max = "", 0
    for url in sources:
        if not url or any(x in url for x in REGISTRES):
            continue
        dom = domaine(url)
        chemin = plain(unquote(urlparse(url).path))
        score = sum(1 for m in mots if m in chemin)              # mot du fait dans le chemin de l'URL
        if dom and dom in preuve.lower():                          # domaine exact cité (« exemple.ch »)
            score += 3
        elif dom and plain(dom.split(".")[0]).replace(" ", "") in plain(preuve).replace(" ", ""):
            score += 2
        if "linkedin" in dom and "linkedin" in preuve.lower():
            score += 2
        if score > score_max:
            meilleur, score_max = url, score
    return meilleur if score_max >= 1 else ""


def faits_sources(d: Donnees, e: dict, contacts: list[dict]) -> tuple[list[dict], dict]:
    """Faits numérotés avec leur source et leur lien : l'agent choisit un id, le script pose le lien (rien d'inventé)."""
    faits, n = [], 0
    for s in sorted(d.signaux.get(e["id"], []) or [], key=lambda s: s.get("date") or "", reverse=True):
        if s.get("famille") == "negatif" or s.get("confiance") == "auto":
            continue
        n += 1
        faits.append({"id": f"S{n}", "date": fmt_date(s.get("date")), "fait": sans_tirets(s.get("resume", "")),
                      "source": libelle_source(s.get("url") or "", e["site"]) or s.get("source", ""),
                      "url": s.get("url") or ""})
    rech = d.recherche.get(e["id"], {}) or {}
    act = rech.get("activite") or e["activite"]
    if act and e["site"]:
        faits.append({"id": "A", "date": "", "fait": act, "source": "site de l'entreprise", "url": e["site"]})
    li = d.linkedin.get(e["id"], {}) or {}
    if li.get("description") and li.get("url"):
        faits.append({"id": "L", "date": "", "fait": clip(li["description"], 500), "source": "page LinkedIn de l'entreprise",
                      "url": li["url"]})
    v = d.logiciel(e["id"])
    if v.get("logiciel") and (v.get("statut") or "").startswith(STATUTS_SURS) and v.get("url"):
        faits.append({"id": "E", "date": "", "fait": f"Logiciel {v['logiciel']} : {sans_tirets(v.get('extrait') or v.get('source') or '')}",
                      "source": libelle_source(v["url"], e["site"]), "url": v["url"]})
    at = rech.get("adoption_tech") or {}
    for i, x in enumerate(at.get("preuves", [])[:3], 1):
        url = lien_preuve(x, (rech.get("sources") or []) + [e["site"]])
        dt = re.search(r"(20\d\d)-(\d\d)-\d\d|\d\d\.(\d\d)\.(20\d\d)", x)
        dt_txt = (f"{dt[2]}.{dt[1]}" if dt[1] else f"{dt[3]}.{dt[4]}") if dt else ""
        faits.append({"id": f"T{i}", "date": dt_txt, "fait": sans_tirets(x),
                      "source": libelle_source(url, e["site"]) if url else "fiche entreprise", "url": url})
    par_contact = {}
    loc = localisation(d, e)
    for c in contacts:
        fc = []
        li_p = c.get("linkedin") or ""
        for j, f in enumerate(faits_personne(d, c, e, loc), 1):
            if f.startswith(("Basé à", "Déjà contacté", "Cité dans un signal")):
                continue
            if f.startswith("Rôle officiel"):
                src, url = "registre du commerce ou site de l'entreprise", ""
            else:
                src, url = "profil LinkedIn", li_p
            fc.append({"id": f"P{j}", "date": "", "fait": f, "source": src, "url": url})
        par_contact[c["contact_id"]] = fc
    return faits, par_contact


def membres_lots(dossier: Path) -> dict:
    """Lot -> (ids d'entreprise, clés de contact) d'après les briefs d'un dossier ; {} sans briefs."""
    out = {}
    for f in sorted(dossier.glob("L*.json")):
        lot = [x for x in lire_json(f, []) or [] if isinstance(x, dict)]
        out[f.name] = ({x.get("id") for x in lot}, {cle_json(c) for x in lot for c in x.get("contacts") or []} - {""})
    return out


def reponses_du_lot(dossier: Path, membres: dict, nom: str):
    """Fichiers L*.json d'un dossier de réponses d'agents, chacun avec les membres de son lot (None : pas de filtre,
    aucun brief). Un fichier sans brief du même nom vient d'un ancien découpage : ignoré, signalé."""
    for f in sorted(dossier.glob("L*.json")):
        if membres and f.name not in membres:
            print(f"! {nom}/{f.name} ignoré : aucun lot {f.stem} dans les briefs (découpage précédent ?)")
            continue
        yield f, membres.get(f.name)


def briefs_sources(d: Donnees) -> None:
    """Lots pour réécrire chaque argument en ouvrant sur sa source (même composition que appels/briefs/L*.json)."""
    _, red_contact = charger_redaction(d.camp)
    base = d.camp.appels / "sources"
    (base / "briefs").mkdir(parents=True, exist_ok=True)
    (base / "redaction").mkdir(parents=True, exist_ok=True)
    par_id = {c["contact_id"]: c for c in d.entree}
    lots = sorted((d.camp.appels / "briefs").glob("L*.json"))
    if not lots:
        raise SystemExit("Aucun brief : lancer d'abord la commande briefs, puis faire rédiger appels/redaction/L*.json.")
    if not red_contact:
        print("! aucune rédaction dans appels/redaction/ : argument_actuel restera vide")
    for f in lots:
        lot = []
        for x in lire_json(f, []) or []:
            e = d.entreprises.get(x["id"])
            if not e:
                continue
            contacts = [d.corriger_nom(par_id[cle_json(c)], e)[0] for c in x["contacts"] if cle_json(c) in par_id]
            faits, par_contact = faits_sources(d, e, contacts)
            lot.append({
                "id": e["id"], "entreprise": x["entreprise"], "pays": x["pays"], "code_pays": x.get("code_pays", ""),
                "lieux_a_eviter": x.get("lieux_a_eviter", []), "logiciels": x.get("logiciels", ""),
                "a_ne_jamais_evoquer": x["a_ne_jamais_evoquer"], "faits_entreprise": faits,
                "contacts": [{"contact_id": c["contact_id"], "prenom": c["prenom"], "nom": c["nom"], "titre": c["titre"],
                              "argument_actuel": (red_contact.get(c["contact_id"]) or {}).get("argument", ""),
                              "faits_personne": par_contact[c["contact_id"]]} for c in contacts],
            })
        (base / "briefs" / f.name).write_text(json.dumps(lot, ensure_ascii=False, indent=1), encoding="utf-8")
        print(f"{f.stem} : {len(lot)} entreprises, {sum(len(x['contacts']) for x in lot)} contacts, "
              f"{sum(len(x['faits_entreprise']) for x in lot)} faits sourcés")
    print(f"-> {base / 'briefs'} ; réponses attendues dans {base / 'redaction'}")


def charger_sources(camp: Campagne) -> dict:
    """contact_id -> {argument, source} : la source est relue dans la fiche à partir de l'id choisi. Une réponse
    n'est reprise que pour un contact du lot de même numéro (appels/sources/briefs)."""
    base = camp.appels / "sources"
    faits = {}                                  # (lot, clé du contact) -> faits numérotés
    for f in sorted((base / "briefs").glob("L*.json")):
        for x in lire_json(f, []) or []:
            for c in x["contacts"]:
                if cle_json(c):
                    faits[(f.name, cle_json(c))] = {ft["id"]: ft for ft in x["faits_entreprise"] + c["faits_personne"]}
    out, hors_lot = {}, 0
    for f, membres in reponses_du_lot(base / "redaction", membres_lots(base / "briefs"), "sources/redaction"):
        for c in lire_json(f, []) or []:
            if not isinstance(c, dict) or not cle_json(c) or (membres and cle_json(c) not in membres[1]):
                hors_lot += 1
                continue
            ft = faits.get((f.name, cle_json(c)), {}).get(c.get("source_id", ""))
            src = ""
            if ft:
                src = ft["source"] + (f", {ft['date']}" if ft["date"] else "") + (f" : {ft['url']}" if ft["url"] else "")
            out[cle_json(c)] = {"argument": c.get("argument", ""), "source": src,
                                "source_id": c.get("source_id", ""), "valide": bool(ft)}
    if hors_lot:
        print(f"! appels/sources/redaction : {hors_lot} réponse(s) ignorée(s), contact absent du lot de même numéro "
              "ou contact_id vide")
    return out


# ---------------------------------------------------------------- assemblage

# Ce que l'appelant lit avant de composer le numéro vient d'abord (contact, argument, pourquoi maintenant), le détail ensuite
COLONNES = ["N°", "Tier", "Entreprise", "Prénom", "Nom", "Poste", "Téléphone", "Email", "Langue",
            "Top argument au téléphone", "Source de l'argument", "Pourquoi maintenant", "Secteur", "Ce qu'ils font", "Taille", "Pays",
            "Canton / département", "Ville", "Logiciels", "Signaux personne", "Signaux direction et actionnariat",
            "Signaux finances", "Signaux activité et projets", "Signaux tech", "Relais à demander",
            "À savoir (ne pas dire)", "Standard", "LinkedIn", "Site web", "ID contact Apollo"]
LARGEURS = {"N°": 5, "Tier": 9, "Entreprise": 22, "Prénom": 12, "Nom": 14, "Poste": 24, "Téléphone": 20,
            "Standard": 16, "Email": 30, "Langue": 13, "Secteur": 22, "Ce qu'ils font": 42, "Taille": 24, "Pays": 10,
            "Canton / département": 18, "Ville": 14, "Signaux personne": 55, "Signaux direction et actionnariat": 55,
            "Signaux finances": 40, "Signaux activité et projets": 50, "Signaux tech": 45, "Logiciels": 32,
            "Pourquoi maintenant": 50, "Top argument au téléphone": 55, "Source de l'argument": 34, "Relais à demander": 45,
            "À savoir (ne pas dire)": 40, "LinkedIn": 30, "Site web": 24, "ID contact Apollo": 26}
COL_RELAIS = ["Entreprise", "N° d'appel du dirigeant", "Prénom", "Nom", "Poste", "Ce qu'il couvre", "Téléphone",
              "Email", "Statut email", "Signaux personne", "LinkedIn", "Ville", "ID contact Apollo"]


def puces(lignes: list[str]) -> str:
    return "\n".join(f"• {x}" for x in lignes if x) if lignes else ""


def charger_redaction(camp: Campagne) -> tuple[dict, dict]:
    """(id d'entreprise -> rédaction, clé du contact -> rédaction), repris seulement du lot de même numéro que le brief
    de l'entreprise et du contact (appels/briefs) ; un contact_id vide n'est jamais repris."""
    par_ent, par_contact, hors_lot = {}, {}, 0
    for f, membres in reponses_du_lot(camp.appels / "redaction", membres_lots(camp.appels / "briefs"), "redaction"):
        ents, cles = membres or (None, None)
        for x in lire_json(f, []) or []:
            if not isinstance(x, dict):
                continue
            if ents is None or x.get("id") in ents:
                par_ent[x.get("id")] = x
            else:
                hors_lot += 1
            for c in x.get("contacts", []):
                k = cle_json(c)
                if k and (cles is None or k in cles):
                    par_contact[k] = c
                else:
                    hors_lot += 1
    if hors_lot:
        print(f"! appels/redaction : {hors_lot} rédaction(s) ignorée(s), entreprise ou contact absent du lot de même "
              "numéro, ou contact_id vide")
    return par_ent, par_contact


# Mots à majuscule permis sans figurer dans la fiche (sigles courants, mots de l'offre, pays)
NOMS_LIBRES = {"ia", "erp", "crm", "mes", "sav", "cao", "pdf", "iso", "mail", "mails", "excel", "linkedin", "vous",
               "votre", "vos", "j", "c", "l", "d", "qu", "n", "s", "r", "it", "europe"}


def controler(d: Donnees, lignes: list[dict]) -> list[str]:
    """Contrôles mécaniques de ce que l'appelant dira : motifs interdits, lieu d'un autre pays, longueur, chiffres et
    noms absents de la fiche de l'entreprise et de l'offre (anti-invention), logiciel nommé sans être confirmé."""
    fiches = {}
    for f in sorted((d.camp.appels / "briefs").glob("L*.json")):
        for x in lire_json(f, []) or []:
            fiches[x["id"]] = json.dumps(x, ensure_ascii=False)
    offre = " ".join(textes_offre(d.cfg) + [d.cfg["offre"].get(k) or "" for k in ("societe", "appelant")])
    libres = NOMS_LIBRES | set(plain(offre).split()) | set(plain(" ".join(NOMS_PAYS.values())).split()) \
        | set(plain(" ".join(w for v in LIEUX.values() for w in v)).split())
    motifs = motifs_interdits(d.cfg)
    surveilles = pays_surveilles(d.cfg)
    sans_sep = lambda t: re.sub(r"(?<=\d)[ '\u2019\u202f\u00a0](?=\d{3}\b)", "", t)   # 7'000, 7 000 -> 7000
    alertes = []
    avec_sources = any(l["Source de l'argument"] for l in lignes)
    for l in lignes:
        qui = f"n°{l['N°']} {l['Prénom']} {l['Nom']} ({l['Entreprise']})"
        arg, eid = l["Top argument au téléphone"], l["_id"]
        if avec_sources and not l["Source de l'argument"]:
            alertes.append(f"{qui} : source de l'argument manquante ou id inconnu")
        if avec_sources and not re.match(r"J'ai (vu|lu|remarqué|noté)", arg):
            alertes.append(f"{qui} : l'argument ne s'ouvre pas sur « J'ai vu »")
        if not arg or arg.startswith("À COMPLÉTER"):
            alertes.append(f"{qui} : argument manquant")
            continue
        for x in interdits_trouves(arg, motifs):
            alertes.append(f"{qui} : motif interdit dans l'argument (« {x} »)")
        for x in lieux_hors_pays(arg, l["_code"], surveilles):
            alertes.append(f"{qui} : « {x} » annoncé à une entreprise hors de ce pays")
        n = len(arg.split())
        if n > 60:
            alertes.append(f"{qui} : argument trop long ({n} mots)")
        fiche = fiches.get(eid, "")
        for nb in re.findall(r"\d+(?:[.,'\u2019 \u202f\u00a0]\d+)*", arg):
            if sans_sep(nb) not in sans_sep(fiche) and sans_sep(nb) not in sans_sep(offre):
                alertes.append(f"{qui} : chiffre « {nb} » absent de la fiche et de l'offre")
        mots_fiche = set(plain(fiche).split())
        for ph in re.split(r"(?<=[.:;?])\s+", arg):
            for mot in re.findall(r"\b[\w'+-]+\b", ph)[1:]:          # le premier mot d'une phrase a sa majuscule
                if re.search(r"[A-ZÀ-Ý]", mot):
                    morceaux = [m for m in plain(mot).split() if m not in libres
                                and m not in plain(l["Source de l'argument"]).replace(" ", "")]   # média tiré du domaine
                    if morceaux and not all(m in mots_fiche for m in morceaux):
                        alertes.append(f"{qui} : nom « {mot} » absent de la fiche (à vérifier)")
        v = d.logiciel(eid)
        for nom in d.cfg["signaux"].get("logiciels") or []:
            if re.search(rf"\b{re.escape(nom)}\b", arg, re.I) and not (
                    nom.lower() in (v.get("logiciel") or "").lower() and (v.get("statut") or "").startswith(STATUTS_SURS)):
                alertes.append(f"{qui} : logiciel « {nom} » nommé sans être confirmé")
        for x in interdits_trouves(l["Pourquoi maintenant"], motifs):
            alertes.append(f"{qui} : motif interdit dans « pourquoi maintenant » (« {x} »)")
        for col in ("Secteur", "Ce qu'ils font", "Taille"):
            if not l[col] or l[col].startswith("À COMPLÉTER"):
                alertes.append(f"{qui} : {col} manquant")
    return alertes


def ecrire_xlsx(lignes: list[dict], relais: list[dict], lisez_moi: list[str], path: Path) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = Workbook()

    def feuille(ws, cols, rows, largeurs, figer):
        ws.append(cols)
        for r in rows:
            ws.append([r.get(c, "") for c in cols])
        entete = PatternFill("solid", fgColor="1F3A5F")
        for cell in ws[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = entete
            cell.alignment = Alignment(wrap_text=True, vertical="center")
        for i, c in enumerate(cols, 1):
            ws.column_dimensions[get_column_letter(i)].width = largeurs.get(c, 18)
        for row in ws.iter_rows(min_row=2):
            hauteur = 1
            for cell, c in zip(row, cols):
                cell.alignment = Alignment(wrap_text=True, vertical="top")
                texte = str(cell.value or "")
                larg = max(largeurs.get(c, 18) - 2, 5)
                hauteur = max(hauteur, sum(max(1, -(-len(x) // larg)) for x in texte.split("\n")))
            ws.row_dimensions[row[0].row].height = min(15 * hauteur + 4, 180)   # au-delà : texte complet dans la cellule
        ws.freeze_panes = figer
        ws.auto_filter.ref = ws.dimensions

    ws = wb.active
    ws.title = "Appels"
    feuille(ws, COLONNES, lignes, LARGEURS, "F2")
    arg_col = COLONNES.index("Top argument au téléphone") + 1
    for row in ws.iter_rows(min_row=2, min_col=arg_col, max_col=arg_col):
        row[0].fill = PatternFill("solid", fgColor="FFF4CC")
        row[0].font = Font(bold=True)
    ws2 = wb.create_sheet("Relais")
    feuille(ws2, COL_RELAIS, relais, {"Entreprise": 22, "N° d'appel du dirigeant": 10, "Prénom": 12, "Nom": 14,
                                      "Poste": 28, "Ce qu'il couvre": 26, "Téléphone": 20, "Email": 30,
                                      "Statut email": 11, "Signaux personne": 60, "LinkedIn": 30, "Ville": 14,
                                      "ID contact Apollo": 26}, "E2")
    ws3 = wb.create_sheet("Lisez-moi")
    for ligne in lisez_moi:
        ws3.append([ligne])
    ws3.column_dimensions["A"].width = 130
    for row in ws3.iter_rows():
        row[0].alignment = Alignment(wrap_text=True, vertical="top")
    ws3["A1"].font = Font(bold=True, size=13)
    wb.save(path)


def lisez_moi(d: Donnees) -> list[str]:
    o, cfg = d.cfg["offre"], d.cfg
    liste = (cfg["apollo"].get("liste_contacts") or "").strip()
    preuves = [p.get("cas") or "preuve" for p in o.get("preuves") or [] if isinstance(p, dict) and p.get("phrase")]
    langues = ", ".join(NOMS_LANGUES.get(l, l) for l in cfg["appels"].get("langues") or []) or "non précisées"
    sig = d.camp.dossier / "donnees" / "signaux"
    jamais = [x for x in o.get("ne_jamais_dire") or [] if x.strip()]
    return [
        f"Fichier d'appel de la campagne {d.camp.nom}" + (f" : liste « {liste} - entrée »" if liste else "")
        + ", dans l'ordre d'appel",
        "",
        "Onglet Appels : une ligne par personne à appeler, dans l'ordre de la colonne N° (d'abord ceux qui parlent une "
        f"langue de l'appelant ({langues}), puis mobile avant standard, puis tier). Onglet Relais : les profils "
        "opérationnels vers lesquels router quand le dirigeant renvoie la balle.",
        "Colonne jaune « Top argument au téléphone » (juste après le contact) : la phrase à placer après la question "
        "ouverte, ou quand la personne demande « et concrètement ? ». Elle s'ouvre sur sa source (« J'ai vu sur votre "
        "page LinkedIn que ... ») ; la colonne « Source de l'argument » donne le lien et la date, à ouvrir avant l'appel.",
        ("Preuves citables : " + ", ".join(preuves) + " ; chacune seulement pour son propre cas d'usage, sans autre "
         "chiffre que le sien.") if preuves else "Aucune preuve dans la configuration : ne citer aucun client ni aucun chiffre.",
        "Signaux : chaque ligne porte sa date (mois.année) et sa source (registre = registre du commerce). « indice » "
        "ou « à vérifier » = pas confirmé : ne pas l'affirmer au prospect, s'en servir pour poser une question.",
        "Logiciels : « confirmé » = l'entreprise le dit elle-même (offre d'emploi, site, LinkedIn) ; « probable » = indice "
        "fort mais indirect ; « indice » = détection d'un outil tiers ou mention indirecte. Inconnu : bonne question "
        "d'ouverture. Taille et chiffre d'affaires : les estimations Apollo sont souvent approximatives (et en USD).",
        "À savoir (ne pas dire) : les signaux négatifs (difficultés, restructuration) ne s'évoquent jamais"
        + (" ; ni : " + " ; ".join(jamais) if jamais else "") + ".",
        "Langue : « autre langue » = la personne parle probablement une autre langue que celles de l'appelant ; "
        "l'argument est écrit dans la langue de rédaction, à adapter.",
        "Lieu : ne jamais situer l'offre dans un autre pays que celui de l'entreprise (colonne Pays).",
        f"Sources : ordre d'appel du {date_fichier(d.chemin_ordre)}, signaux du {date_fichier(sig / 'signaux.json')} "
        f"(registre, LinkedIn, sites, presse, recherches contre-vérifiées), logiciels relus le "
        f"{date_fichier(sig / 'logiciels.json')}, parcours des personnes d'après Apollo (LinkedIn). Date de référence : "
        f"{REF.strftime('%d.%m.%Y')}.",
    ]


# Colonnes ajoutées au fichier d'ordre d'appel (--maj-ordre), juste avant la colonne angle
COLONNES_ORDRE = {"top_argument": "Top argument au téléphone", "source_argument": "Source de l'argument",
                  "pourquoi_maintenant": "Pourquoi maintenant", "secteur": "Secteur", "ce_qu_ils_font": "Ce qu'ils font",
                  "taille": "Taille", "canton_departement": "Canton / département", "standard": "Standard",
                  "logiciels": "Logiciels", "signaux_personne": "Signaux personne",
                  "signaux_direction": "Signaux direction et actionnariat", "signaux_finances": "Signaux finances",
                  "signaux_activite": "Signaux activité et projets", "signaux_tech": "Signaux tech",
                  "relais_a_demander": "Relais à demander", "a_savoir_ne_pas_dire": "À savoir (ne pas dire)"}


def completer_ordre(chemin: Path, lignes: list[dict]) -> int:
    """Écrit les colonnes du fichier dans le fichier d'ordre d'appel (même ordre de lignes), prénom et nom corrigés.
    Rattachement par la clé du contact (id Apollo ou id local), écrite dans contact_id si elle manquait : un prénom
    corrigé ne la change donc plus. Idempotent. À relancer après ordre_appel.py (qui réécrit ce fichier) et
    scripts_appel.py --maj-ordre."""
    contacts = lire_csv(chemin)
    par_id = {l["_cle"]: l for l in lignes if l.get("_cle")}
    base = [k for k in contacts[0] if k not in COLONNES_ORDRE] if contacts else []
    if contacts and "contact_id" not in base:
        base.append("contact_id")
    i = base.index("angle") if "angle" in base else len(base)
    champs = base[:i] + list(COLONNES_ORDRE) + base[i:]
    for x in contacts:
        cle = cle_contact(x)
        l = par_id.get(cle) if cle else None
        if not l:
            continue
        if not (x.get("contact_id") or "").strip():
            x["contact_id"] = cle
        x["prenom"], x["nom"] = l["Prénom"], l["Nom"]
        for col, src in COLONNES_ORDRE.items():
            x[col] = l[src]
    ecrire_csv(chemin, champs, contacts)
    return len(contacts)


def assembler(d: Donnees, maj_ordre: bool = False) -> None:
    red_ent, red_contact = charger_redaction(d.camp)
    sources = charger_sources(d.camp)
    lignes, rang_entreprise = [], {}
    for c in d.entree:
        e = d.entreprises[c["_id"]]
        loc = localisation(d, e)
        r_e = red_ent.get(e["id"], {}) or {}
        r_c = red_contact.get(c["contact_id"], {}) or {}
        rang_entreprise.setdefault(e["id"], c["priorite"])
        ui = d.ui.get(c["contact_id"], {})
        if ui.get("Last Name"):              # noms d'Apollo d'origine : la liste a pu être corrigée par une passe précédente
            c = {**c, "prenom": ui.get("First Name", ""), "nom": ui["Last Name"]}
        c, note_nom = d.corriger_nom(c, e)
        a_savoir = [note_nom] + [f"Ne pas évoquer : {x}" for x in signaux(d, e, ("negatif",))]
        a_savoir.append((d.corrections.get(e["id"], {}) or {}).get("a_savoir", ""))
        if c.get("a_eviter"):
            a_savoir.append(f"À éviter : {c['a_eviter']}")
        if loc["note"]:
            a_savoir.append(f"Lieu : {loc['note']}.")
        lieux = lieux_a_eviter(loc["code"], d.cfg)
        if lieux:
            a_savoir.append("Ne pas dire à cette entreprise : " + ", ".join(f"« {x} »" for x in lieux) + ".")
        if not c["id_apollo"]:
            a_savoir.append("Contact absent d'Apollo : noter l'appel à la main (ou créer le contact).")
        lignes.append({
            "_id": e["id"], "_cle": c["contact_id"], "_code": loc["code"],
            "N°": int(c["priorite"]) if str(c["priorite"]).isdigit() else c["priorite"],
            "Tier": e["tier"] + (" (fenêtre ouverte)" if c.get("fenetre_opportunite") == "oui" else ""),
            "Entreprise": e["nom"] or c["entreprise"], "Prénom": c["prenom"], "Nom": c["nom"],
            "Poste": (c["titre"] or ui.get("Title") or (d.personne(c) or {}).get("title")
                      or (role_officiel(d, c, e) or {}).get("role", "").split(";")[0].strip() or "inconnu"),
            "Téléphone": telephone_txt(c), "Standard": standard(d, e, loc["code"]), "Email": c["email"] or "aucun",
            "Langue": langue_txt(c, d.langues),
            "Secteur": r_e.get("secteur") or "À COMPLÉTER", "Ce qu'ils font": r_e.get("description") or "À COMPLÉTER",
            "Taille": r_e.get("taille") or "À COMPLÉTER",
            "Pays": loc["pays"], "Canton / département": loc["region"], "Ville": loc["ville"],
            "Signaux personne": puces(faits_personne(d, c, e, loc)),
            "Signaux direction et actionnariat": puces(direction(d, e)) or "aucun signal trouvé",
            "Signaux finances": puces(finances(d, e)) or "aucun signal trouvé",
            "Signaux activité et projets": puces(signaux(d, e, ("croissance", "projet"))) or "aucun signal trouvé",
            "Signaux tech": puces(tech(d, e)) or "aucun signal trouvé",
            "Logiciels": logiciels_txt(d, e),
            "Pourquoi maintenant": pourquoi_maintenant(d, e, r_e.get("pourquoi_maintenant") or ""),
            "Top argument au téléphone": (sources.get(c["contact_id"]) or r_c).get("argument") or "À COMPLÉTER",
            "Source de l'argument": (sources.get(c["contact_id"]) or {}).get("source", ""),
            "Relais à demander": relais_txt(d, e["id"]),
            "À savoir (ne pas dire)": puces([x for x in a_savoir if x]),
            "LinkedIn": c["linkedin"], "Site web": e["site"], "ID contact Apollo": c["id_apollo"],
        })
    relais = []
    for cle, rs in d.relais.items():
        e = d.entreprises.get(cle)
        for r in rs:
            r, _ = d.corriger_nom(r, e)
            faits = faits_personne(d, {**r, "deja_contacte": ""}, e, localisation(d, e)) if e else []
            n_dir = rang_entreprise.get(cle, "")
            relais.append({"Entreprise": (e or {}).get("nom") or r["entreprise"],
                           "N° d'appel du dirigeant": int(n_dir) if str(n_dir).isdigit() else n_dir,
                           "Prénom": r["prenom"], "Nom": r["nom"], "Poste": r["titre"], "Ce qu'il couvre": cibles_txt(r),
                           "Téléphone": fmt_tel(r.get("mobile")) or (f"standard {fmt_tel(r['telephone_standard'])}"
                                                                     if r.get("telephone_standard") else "aucun"),
                           "Email": r["email"] or "aucun", "Statut email": r["statut_email"],
                           "Signaux personne": puces(faits),
                           "LinkedIn": r["linkedin"], "Ville": r["ville"], "ID contact Apollo": r["contact_id"]})
    relais.sort(key=lambda r: (r["N° d'appel du dirigeant"] == "", str(r["N° d'appel du dirigeant"]).zfill(6), r["Entreprise"]))
    alertes = controler(d, lignes)
    xlsx, csv_ = d.camp.appels / "fichier-appel.xlsx", d.camp.appels / "fichier-appel.csv"
    ecrire_xlsx(lignes, relais, lisez_moi(d), xlsx)
    ecrire_csv(csv_, COLONNES, lignes)
    print(f"{len(lignes)} contacts, {len({l['_id'] for l in lignes})} entreprises, {len(relais)} relais")
    if maj_ordre:
        n = completer_ordre(d.chemin_ordre, lignes)
        print(f"-> {n} lignes complétées dans {d.chemin_ordre} ({len(COLONNES_ORDRE)} colonnes)")
    print(f"-> {xlsx}\n-> {csv_}")
    print(f"{len(alertes)} alerte(s)" + "".join(f"\n  - {x}" for x in alertes))


def main() -> None:
    ap = argparse.ArgumentParser(description="Fichier d'appel du commercial (étape 9, aucun appel réseau).")
    ap.add_argument("commande", choices=["briefs", "sources", "assembler"],
                    help="briefs (fiches pour les agents), sources (faits numérotés), assembler (xlsx et csv)")
    argument_campagne(ap)
    ap.add_argument("--lots", type=int, default=4, help="briefs : nombre de lots, un agent par lot (défaut 4)")
    ap.add_argument("--ordre", help="CSV d'ordre d'appel (défaut : appels/ordre-appel.csv)")
    ap.add_argument("--relais", help="CSV des relais (défaut : export « - relais » de la liste, sinon contacts-crees.csv)")
    ap.add_argument("--export", help="facultatif : export CSV de la liste « entrée » fait depuis l'interface Apollo")
    ap.add_argument("--maj-ordre", action="store_true",
                    help="assembler : écrire aussi les colonnes dans le fichier d'ordre d'appel (non par défaut)")
    a = ap.parse_args()
    camp = Campagne(a.campagne)
    d = Donnees(camp, a)
    if a.commande == "briefs":
        briefs(d, a.lots)
    elif a.commande == "sources":
        briefs_sources(d)
    else:
        assembler(d, a.maj_ordre)


if __name__ == "__main__":
    main()
